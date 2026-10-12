import os
import re
from typing import List, Dict, Any
from pathlib import Path
import pymupdf as fitz  # PyMuPDF
import docx  # python-docx

class DocumentParsingError(Exception):
    pass

class DocumentParser:
    @staticmethod
    def clean_text(text: str) -> str:
        """Clean extracted text while preserving semantic structure."""
        if not text:
            return ""
        # Normalize carriage returns and weird unicode whitespace
        text = text.replace("\r\n", "\n").replace("\r", "\n")
        # Replace non-breaking spaces
        text = text.replace("\u00a0", " ")
        # Reduce more than two consecutive newlines to two
        text = re.sub(r"\n{3,}", "\n\n", text)
        return text.strip()

    @classmethod
    def parse_pdf(cls, file_path: str) -> Dict[str, Any]:
        """
        Extract text from PDF using PyMuPDF (fitz), preserving page numbers.
        """
        try:
            doc = fitz.open(file_path)
            total_pages = len(doc)
            pages: List[Dict[str, Any]] = []

            for page_index in range(total_pages):
                page = doc[page_index]
                page_number = page_index + 1
                text = page.get_text("text")
                cleaned = cls.clean_text(text)
                
                # Only include pages that have extracted text
                if cleaned:
                    pages.append({
                        "page_number": page_number,
                        "text": cleaned,
                        "char_count": len(cleaned)
                    })

            doc.close()

            if not pages:
                raise DocumentParsingError(
                    "The PDF document does not contain extractable text. It may be scanned or image-only."
                )

            return {
                "file_type": "pdf",
                "total_pages": total_pages,
                "pages": pages
            }

        except fitz.FileDataError as e:
            raise DocumentParsingError(f"Corrupt or invalid PDF file: {str(e)}")
        except Exception as e:
            if isinstance(e, DocumentParsingError):
                raise
            raise DocumentParsingError(f"Failed to parse PDF document: {str(e)}")

    @classmethod
    def parse_docx(cls, file_path: str) -> Dict[str, Any]:
        """
        Extract text from Word document using python-docx.
        Segments paragraphs and tables into logical pages.
        """
        try:
            doc = docx.Document(file_path)
            pages: List[Dict[str, Any]] = []
            current_page_text: List[str] = []
            current_page_number = 1
            approx_page_char_target = 1800  # standard academic page ~300 words (~1800 chars)

            # Process paragraphs
            for para in doc.paragraphs:
                text = para.text.strip()
                if not text:
                    continue

                # Check if paragraph contains a hard page break
                has_page_break = any(
                    run._r.xpath(".//w:br[@w:type='page']") for run in para.runs
                )

                current_page_text.append(text)

                total_chars = sum(len(t) for t in current_page_text)
                if has_page_break or total_chars >= approx_page_char_target:
                    combined = cls.clean_text("\n\n".join(current_page_text))
                    if combined:
                        pages.append({
                            "page_number": current_page_number,
                            "text": combined,
                            "char_count": len(combined)
                        })
                        current_page_number += 1
                        current_page_text = []

            # Process tables (academic papers often have tables)
            for table in doc.tables:
                table_lines = []
                for row in table.rows:
                    row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    row_text = " | ".join(row_cells)
                    if row_text:
                        table_lines.append(row_text)
                
                if table_lines:
                    current_page_text.append("\n[Table Data]:\n" + "\n".join(table_lines))

            # Flush remaining text
            if current_page_text:
                combined = cls.clean_text("\n\n".join(current_page_text))
                if combined:
                    pages.append({
                        "page_number": current_page_number,
                        "text": combined,
                        "char_count": len(combined)
                    })

            if not pages:
                raise DocumentParsingError(
                    "The DOCX document does not contain any extractable text content."
                )

            return {
                "file_type": "docx",
                "total_pages": len(pages),
                "pages": pages
            }

        except Exception as e:
            if isinstance(e, DocumentParsingError):
                raise
            raise DocumentParsingError(f"Failed to parse DOCX document: {str(e)}")

    @classmethod
    def parse_document(cls, file_path: str) -> Dict[str, Any]:
        """
        Auto-detects document format by extension and parses text with page metadata.
        """
        ext = Path(file_path).suffix.lower()
        if ext == ".pdf":
            return cls.parse_pdf(file_path)
        elif ext in [".docx", ".doc"]:
            return cls.parse_docx(file_path)
        else:
            raise DocumentParsingError(
                f"Unsupported file format '{ext}'. Supported formats: PDF, DOCX."
            )
