"""
Script to generate sample academic PDF and DOCX files for immediate testing.
"""
import os
from pathlib import Path

def create_sample_pdf(output_path: str):
    import fitz  # PyMuPDF
    
    doc = fitz.open()
    
    # Page 1: Title, Abstract, Introduction
    page1 = doc.new_page(width=595, height=842) # A4
    text_p1 = (
        "A Retrieval-Augmented Generation Framework for Academic QA\n\n"
        "Dr. Eleanor Vance, Prof. Marcus Chen\n"
        "Department of Computer Science & Information Engineering\n\n"
        "Abstract\n"
        "Academic literature retrieval poses distinct challenges due to domain-specific terminology, "
        "dense mathematical formulations, and the critical requirement of factual grounding. "
        "Large Language Models (LLMs) frequently hallucinate facts when queried on non-parametric knowledge. "
        "In this study, we propose an intelligent question-answering architecture based on Retrieval-Augmented "
        "Generation (RAG). By integrating PyMuPDF text extraction, semantic chunking with overlap, dense vector "
        "indexing via ChromaDB, and context-conditioned Gemini generation, our system achieves a 94.2% factual "
        "precision score across 1,200 peer-reviewed computer science papers.\n\n"
        "1. Introduction\n"
        "Scholarly research demands precise evidence extraction. Traditional keyword-based retrieval mechanisms "
        "(such as BM25) often fail to capture semantic nuances, synonymy, and cross-section conceptual relations. "
        "Conversely, generative pre-trained transformers possess vast linguistic capabilities but suffer from catastrophic "
        "hallucinations and temporal knowledge cutoffs. The fusion of dense vector retrieval with instruction-tuned "
        "generative models provides a mathematically sound solution for evidence-backed question answering.\n\n"
        "Page 1"
    )
    rect = fitz.Rect(50, 50, 545, 792)
    page1.insert_textbox(rect, text_p1, fontsize=11, fontname="helv", color=(0.1, 0.1, 0.1))
    
    # Page 2: Methodology & System Architecture
    page2 = doc.new_page(width=595, height=842)
    text_p2 = (
        "2. System Architecture and Methodology\n\n"
        "Our RAG framework comprises four principal modular stages:\n\n"
        "2.1 Document Ingestion and Page-Aware Extraction\n"
        "Documents in PDF and DOCX formats undergo structural decomposition. For PDF files, PyMuPDF decomposes "
        "the document stream into granular page blocks, extracting raw UTF-8 text streams while retaining sequential "
        "page numbers. Word documents are processed using python-docx with paragraph and table extraction.\n\n"
        "2.2 Semantic Chunking and Overlap Strategies\n"
        "Extracted text is partitioned into overlapping windows of 750 characters with a 150-character sliding buffer. "
        "This boundary preservation prevents truncation of academic arguments spanning sentence boundaries.\n\n"
        "2.3 Dense Vector Indexing in ChromaDB\n"
        "Each chunk is encoded into a 384-dimensional dense semantic embedding using the all-MiniLM-L6-v2 transformer. "
        "Embeddings and metadata (document ID, filename, page index) are ingested into ChromaDB HNSW indexes with cosine distance metric.\n\n"
        "2.4 Context-Conditioned Gemini Generation\n"
        "At query time, the student's question is vectorized. Cosine similarity retrieves top-k relevant chunks. A constrained "
        "grounding prompt instructs Gemini to formulate an answer strictly supported by the retrieved passages, citing "
        "the source document and page number.\n\n"
        "Page 2"
    )
    rect2 = fitz.Rect(50, 50, 545, 792)
    page2.insert_textbox(rect2, text_p2, fontsize=11, fontname="helv", color=(0.1, 0.1, 0.1))
    
    # Page 3: Experimental Results and Conclusions
    page3 = doc.new_page(width=595, height=842)
    text_p3 = (
        "3. Experimental Results and Discussion\n\n"
        "We evaluated the framework on a benchmark dataset of 500 academic research inquiries.\n\n"
        "Table 1: Comparative Evaluation of QA Frameworks\n"
        "Metric                 | Pure Gemini LLM | BM25 Retrieval | Proposed RAG System\n"
        "Faithfulness / Grounding | 61.4%           | 78.2%          | 94.2%\n"
        "Answer Relevance       | 82.0%           | 71.5%          | 96.8%\n"
        "Hallucination Rate     | 38.6%           | 14.1%          | 2.8%\n"
        "Mean Latency (sec)     | 1.2s            | 0.4s           | 1.8s\n\n"
        "The empirical findings demonstrate that dense semantic retrieval reduces hallucination by over 92% compared to "
        "unaugmented generation. Strict prompt guardrails ensure that when queries fall outside the document scope, the "
        "system reliably indicates insufficiency rather than inventing specious claims.\n\n"
        "4. Conclusion and Future Directions\n"
        "This research establishes an effective, lightweight, and scalable academic question answering framework. "
        "Future enhancements include cross-encoder re-ranking and multi-modal equation parsing.\n\n"
        "Page 3"
    )
    rect3 = fitz.Rect(50, 50, 545, 792)
    page3.insert_textbox(rect3, text_p3, fontsize=11, fontname="helv", color=(0.1, 0.1, 0.1))

    doc.save(output_path)
    doc.close()
    print(f"Sample PDF created at {output_path}")

def create_sample_docx(output_path: str):
    import docx
    from docx.shared import Inches, Pt, RGBColor
    
    doc = docx.Document()
    
    # Title
    title = doc.add_heading("Research Proposal: Evaluating Hallucination Mitigation in Academic RAG", level=0)
    
    doc.add_paragraph("Principal Investigator: Dr. Marcus Chen, Stanford AI Lab")
    doc.add_paragraph("Grant Category: Intelligent Educational Technologies\n")
    
    # Section 1
    doc.add_heading("1. Project Rationale & Objectives", level=1)
    doc.add_paragraph(
        "Modern university students grapple with information overload when surveying vast academic corpora. "
        "Generative AI tools promise automated literature synthesis, but their tendency to fabricate citations "
        "and misattribute theoretical findings poses a severe hazard to academic rigor. This project aims to deploy "
        "an open-access, browser-based RAG application that guarantees strict provenance tracking."
    )
    
    # Section 2
    doc.add_heading("2. Technical Milestones and Deliverables", level=1)
    doc.add_paragraph(
        "Milestone 1: Multi-format parsing pipeline supporting PDF and Word (.docx) documents with page number retention.\n"
        "Milestone 2: Sub-second vector retrieval using ChromaDB persistence with cosine indexing.\n"
        "Milestone 3: Strict grounding prompts conditioned on Google Gemini API to eliminate hallucinated answers.\n"
        "Milestone 4: Cloud deployment on Vercel and Render for universal student access."
    )
    
    # Section 3
    doc.add_heading("3. Expected Outcomes and Benchmarking", level=1)
    table = doc.add_table(rows=1, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Phase"
    hdr_cells[1].text = "Objective"
    hdr_cells[2].text = "Target Accuracy"
    
    row1 = table.add_row().cells
    row1[0].text = "Phase A"
    row1[1].text = "Document Extraction & Chunking"
    row1[2].text = "99.5% text retention"
    
    row2 = table.add_row().cells
    row2[0].text = "Phase B"
    row2[1].text = "Vector Similarity Retrieval"
    row2[2].text = "Top-4 Recall > 92%"

    row3 = table.add_row().cells
    row3[0].text = "Phase C"
    row3[1].text = "Grounded Synthesis"
    row3[2].text = "Hallucination < 3%"

    doc.add_page_break()
    doc.add_heading("4. Budget and Resource Allocation", level=1)
    doc.add_paragraph(
        "Total requested funding is $45,000 distributed over cloud computing credits, vector storage bandwidth, "
        "and graduate researcher compensation. The application will be released open-source under the MIT license."
    )

    doc.save(output_path)
    print(f"Sample DOCX created at {output_path}")

if __name__ == "__main__":
    out_dir = Path(__file__).parent / "sample_docs"
    out_dir.mkdir(exist_ok=True)
    create_sample_pdf(str(out_dir / "sample_academic_paper.pdf"))
    create_sample_docx(str(out_dir / "sample_research_proposal.docx"))
